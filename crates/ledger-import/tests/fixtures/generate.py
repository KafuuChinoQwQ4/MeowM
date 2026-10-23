"""Regenerate synthetic Excel fixtures (pip install openpyxl xlwt). No real bills."""

from datetime import datetime
from pathlib import Path

import openpyxl
import xlwt

HERE = Path(__file__).resolve().parent

book = openpyxl.Workbook()
book.active.title = "说明"
book.active.append(["合成测试账单，不含真实个人数据"])
sheet = book.create_sheet("微信账单")
sheet.cell(2, 1, "微信支付账单明细")
sheet.append(["交易时间", "交易类型", "交易对方", "商品", "收/支", "金额(元)", "交易单号"])
sheet.append([])
sheet.append([datetime(2026, 10, 2, 8, 30), "消费", "店,铺", '早餐"套餐"\n饮品', "支出", 12.30, "00012345678901234567890"])
sheet.append([datetime(2026, 10, 3), "收入", "测试公司", "工资", "收入", 0.29, "wx-income"])
sheet.append([])
sheet.append(["invalid-date", "消费", "商户", "", "支出", 1, "wx-invalid"])
for row in (5, 6):
    sheet.cell(row, 1).number_format = "yyyy-mm-dd hh:mm:ss"
    sheet.cell(row, 6).number_format = "0.00"
book.save(HERE / "wechat.xlsx")

date_style = xlwt.easyxf(num_format_str="YYYY-MM-DD HH:MM:SS")
money_style = xlwt.easyxf(num_format_str="0.00")
for filename, epoch_1904, headers, rows in [
    (
        "alipay.xls",
        False,
        ["交易时间", "交易对方", "商品说明", "收/支", "交易金额", "交易订单号"],
        [
            [datetime(2026, 10, 2, 9, 10), "商户", "午餐", "支出", 0.29, "00098765432109876543210"],
            [datetime(2026, 10, 3), "商户", "退货", "退款", 5, "ali-refund"],
        ],
    ),
    (
        "bank-1904.xls",
        True,
        ["交易日期", "摘要", "借贷标记", "交易金额", "流水号"],
        [[datetime(2026, 10, 3), "超市", "借", -100.01, "00000000000000000012345"]],
    ),
]:
    book = xlwt.Workbook()
    book.dates_1904 = epoch_1904
    sheet = book.add_sheet("账单")
    for row_number, row in enumerate([headers, *rows]):
        for column, value in enumerate(row):
            style = date_style if isinstance(value, datetime) else money_style if isinstance(value, (int, float)) else xlwt.Style.default_style
            sheet.write(row_number, column, value, style)
    book.save(str(HERE / filename))

# Calamine's XLS reader allocates a dense range for large sparse extents.
# This small workbook exercises the preflight limit before that allocation.
book = xlwt.Workbook()
sheet = book.add_sheet("账单")
for column, value in enumerate(["日期", "金额", "收支", "流水号"]):
    sheet.write(0, column, value)
sheet.write(1, 0, "2026-10-03")
sheet.write(1, 1, 1)
sheet.write(1, 2, "支出")
sheet.write(1, 3, "ordinary")
sheet.write(65535, 255, "out-of-bounds")
book.save(str(HERE / "sparse-large.xls"))
